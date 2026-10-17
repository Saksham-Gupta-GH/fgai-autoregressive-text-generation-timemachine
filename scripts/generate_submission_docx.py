import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_docx_hyperlink(paragraph, url, text, color="1155CC", underline=True):
    """Inserts an active clickable hyperlink into a python-docx paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def build_submission_docx(
    student_name="Saksham Gupta",
    reg_num="230911186",
    repo_url="https://github.com/Saksham-Gupta-GH/fgai-autoregressive-text-generation-timemachine",
    output_filename="SakshamGupta_230911186_FGAI_Assignment2.docx"
):
    doc = Document()
    
    # Standard 1.0-inch margins matching Google Docs default
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10.5)
        
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("Foundations of Generative AI (ICT4403)")
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)
    
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub.add_run("Assignment 2 - Submission Details")
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(13)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    
    # Metadata Table with clean Table Grid borders
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    meta_rows = [
        ("Student name", student_name, None),
        ("Registration number", reg_num, None),
        ("Assignment title", "Foundations of Generative AI (ICT4403) - Assignment 2", None),
        ("Title of the developed autoregressive text-generation system", "Comparative Word-Level Autoregressive Text Generation on H. G. Wells' The Time Machine using Simple RNN, Bidirectional RNN, LSTM, and GRU", None),
        ("Dataset name", "The Time Machine by H. G. Wells (Project Gutenberg eBook #35)", None),
        ("Dataset link", "", "https://www.gutenberg.org/files/35/35-0.txt"),
        ("GitHub repository link", "", repo_url)
    ]
    
    for i, (k, text_prefix, link_url) in enumerate(meta_rows):
        row = table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.6)
        c1.width = Inches(4.2)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(5)
        p0.paragraph_format.space_after = Pt(5)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(5)
        p1.paragraph_format.space_after = Pt(5)
        if text_prefix:
            r1 = p1.add_run(text_prefix)
            r1.font.size = Pt(10.5)
        if link_url:
            add_docx_hyperlink(p1, link_url, link_url)
            
    output_path = os.path.join(os.getcwd(), output_filename)
    doc.save(output_path)
    print(f"Generated Word document at: {output_path}")
    return output_path


if __name__ == "__main__":
    build_submission_docx()
